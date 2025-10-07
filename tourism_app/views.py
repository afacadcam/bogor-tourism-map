# tourism_app/views.py
import pandas as pd
import openpyxl
import xlrd
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import JsonResponse
from django.views.generic import ListView, CreateView, UpdateView, DeleteView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.contrib.gis.geos import Point
from .models import TourismPlace, TourismCategory, DataImport
from .forms import FileImportForm, TourismPlaceForm
from django.urls import reverse_lazy
from openpyxl.utils import get_column_letter

import logging
import csv
from django.http import HttpResponse

logger = logging.getLogger(__name__)

class FileImportView(LoginRequiredMixin, CreateView):
    model = DataImport
    form_class = FileImportForm
    template_name = 'tourism/import_file.html'
    success_url = '/tourism/import-history/'
    
    def form_valid(self, form):
        form.instance.created_by = self.request.user
        
        # Determine file type
        file_name = form.instance.file.name.lower()
        if file_name.endswith('.csv'):
            form.instance.file_type = 'csv'
        elif file_name.endswith('.xlsx'):
            form.instance.file_type = 'xlsx'
        elif file_name.endswith('.xls'):
            form.instance.file_type = 'xls'
        
        response = super().form_valid(form)
        
        # Process file immediately
        self.process_import_file(self.object)
        
        messages.success(self.request, 'File berhasil diunggah dan diproses!')
        return response
    
    def process_import_file(self, import_obj):
        """Process uploaded file and create TourismPlace objects"""
        try:
            import_obj.status = 'processing'
            import_obj.save()
            
            # Read file based on type
            if import_obj.file_type == 'csv':
                df = pd.read_csv(import_obj.file.path)
            elif import_obj.file_type in ['xlsx', 'xls']:
                engine = 'openpyxl' if import_obj.file_type == 'xlsx' else 'xlrd'
                df = pd.read_excel(import_obj.file.path, engine=engine)
            
            import_obj.total_rows = len(df)
            import_obj.save()
            
            # Expected columns mapping
            column_mapping = {
                'nama': 'name',
                'name': 'name',
                'deskripsi': 'description',
                'description': 'description',
                'kategori': 'category',
                'category': 'category',
                'alamat': 'address',
                'address': 'address',
                'latitude': 'latitude',
                'lat': 'latitude',
                'longitude': 'longitude',
                'lng': 'longitude',
                'lon': 'longitude',
                'telepon': 'phone',
                'phone': 'phone',
                'website': 'website',
                'email': 'email',
                'jam_buka': 'opening_hours',
                'opening_hours': 'opening_hours',
                'harga_tiket': 'ticket_price',
                'ticket_price': 'ticket_price',
                'rating': 'rating'
            }
            
            # Normalize column names
            df.columns = df.columns.str.lower().str.strip()
            df = df.rename(columns=column_mapping)
            
            # Validate required columns
            required_columns = ['name', 'latitude', 'longitude']
            missing_columns = [col for col in required_columns if col not in df.columns]
            
            if missing_columns:
                raise ValueError(f"Missing required columns: {', '.join(missing_columns)}")
            
            processed_count = 0
            error_count = 0
            error_logs = []
            
            for index, row in df.iterrows():
                try:
                    # Get or create category
                    category_name = row.get('category', 'Umum')
                    category, _ = TourismCategory.objects.get_or_create(
                        name=category_name,
                        defaults={'description': f'Kategori {category_name}'}
                    )
                    
                    # Create Point geometry
                    latitude = float(row['latitude'])
                    longitude = float(row['longitude'])
                    location = Point(longitude, latitude, srid=4326)
                    
                    # Create TourismPlace
                    place, created = TourismPlace.objects.get_or_create(
                        name=row['name'],
                        defaults={
                            'description': row.get('description', ''),
                            'category': category,
                            'location': location,
                            'address': row.get('address', ''),
                            'phone': row.get('phone', ''),
                            'website': row.get('website', ''),
                            'email': row.get('email', ''),
                            'opening_hours': row.get('opening_hours', ''),
                            'ticket_price': float(row.get('ticket_price', 0) or 0),
                            'rating': float(row.get('rating', 0) or 0),
                            'created_by': import_obj.created_by
                        }
                    )
                    
                    if created:
                        processed_count += 1
                    else:
                        # Update existing place
                        for field in ['description', 'address', 'phone', 'website', 'email', 
                                    'opening_hours', 'ticket_price', 'rating']:
                            if field in row and pd.notna(row[field]) and row[field] != '':
                                setattr(place, field, row[field])
                        place.location = location
                        place.category = category
                        place.save()
                        processed_count += 1
                
                except Exception as e:
                    error_count += 1
                    error_logs.append(f"Row {index + 2}: {str(e)}")
                    logger.error(f"Error processing row {index + 2}: {e}")
            
            # Update import status
            import_obj.processed_rows = processed_count
            import_obj.error_rows = error_count
            import_obj.error_log = '\n'.join(error_logs)
            import_obj.status = 'completed'
            import_obj.save()
            
        except Exception as e:
            import_obj.status = 'failed'
            import_obj.error_log = str(e)
            import_obj.save()
            logger.error(f"Import failed: {e}")

class ImportHistoryView(LoginRequiredMixin, ListView):
    model = DataImport
    template_name = 'tourism/import_history.html'
    context_object_name = 'imports'
    paginate_by = 20
    
    def get_queryset(self):
        return DataImport.objects.filter(created_by=self.request.user)

class AddPlaceView(LoginRequiredMixin, CreateView):
    model = TourismPlace
    form_class = TourismPlaceForm
    template_name = 'tourism/add_place.html'
    success_url = reverse_lazy('tourism_app:place_list') # Arahkan ke daftar tempat setelah berhasil

    def form_valid(self, form):
        # set created_by dll jika perlu
        if hasattr(form.instance, 'created_by'):
            form.instance.created_by = self.request.user if self.request.user.is_authenticated else None
        messages.success(self.request, 'Tempat wisata berhasil ditambahkan!')
        return super().form_valid(form)
    
    def form_invalid(self, form):
        messages.error(self.request, 'Gagal menyimpan. Periksa kembali data yang diisi.')
        return super().form_invalid(form)


def export_sample_template(request):
    """Export sample CSV template for data import"""
    
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="template_import_pariwisata.csv"'
    
    writer = csv.writer(response)
    # Write header
    writer.writerow([
        'nama', 'deskripsi', 'kategori', 'alamat', 'latitude', 'longitude',
        'telepon', 'website', 'email', 'jam_buka', 'harga_tiket', 'rating'
    ])
    
    # Write sample data
    writer.writerow([
        'Kebun Raya Bogor',
        'Kebun raya tertua di Indonesia dengan koleksi tanaman yang lengkap',
        'Wisata Alam',
        'Jl. Ir. H. Juanda No.13, Paledang, Kecamatan Bogor Tengah, Kota Bogor, Jawa Barat',
        '-6.5972',
        '106.7997',
        '(0251) 8311362',
        'https://krbogor.lipi.go.id/',
        'info@krbogor.lipi.go.id',
        '08:00 - 17:00',
        '20000',
        '4.5'
    ])
    
    return response

def export_places_csv(request):
    response = HttpResponse(
        content_type='text/csv',
        headers={'Content-Disposition': 'attachment; filename="daftar_tempat_wisata.csv"'},
    )
    writer = csv.writer(response)
    # Tulis header
    writer.writerow([
        'No', 'Nama', 'Deskripsi', 'Kategori', 'Alamat', 'Latitude', 'Longitude',
        'Telepon', 'Website', 'Email', 'Jam Buka', 'Harga Tiket', 'Rating'
    ])
    for i, place in enumerate(TourismPlace.objects.all(), start=1):
        writer.writerow([
            i,
            place.name,
            place.description,
            place.category.name if place.category else '',
            place.address,
            place.latitude,
            place.longitude,
            place.phone if place.phone else '',
            place.website if place.website else '',
            place.email if place.email else '',
            place.opening_hours,
            place.ticket_price if place.ticket_price else '',
            place.rating,
        ])
    return response

def export_places_xls(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tempat Wisata"
    ws.append([
        'No', 'Nama', 'Deskripsi', 'Kategori', 'Alamat', 'Latitude', 'Longitude',
        'Telepon', 'Website', 'Email', 'Jam Buka', 'Harga Tiket', 'Rating'
    ])
    for i, place in enumerate(TourismPlace.objects.all(), start=1):
        ws.append([
            i,
            place.name,
            place.description,
            place.category.name if place.category else '',
            place.address,
            place.latitude,
            place.longitude,
            place.phone if place.phone else '',
            place.website if place.website else '',
            place.email if place.email else '',
            place.opening_hours,
            place.ticket_price if place.ticket_price else '',
            place.rating,
        ])

    # === AUTO-FIT LEBAR KOLOM ===
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                cell_value = str(cell.value)
            except:
                cell_value = ""
            if cell_value:
                max_length = max(max_length, len(cell_value))
        # Tambah sedikit padding agar tidak terlalu rapat
        adjusted_width = max_length + 2
        ws.column_dimensions[col_letter].width = adjusted_width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=daftar_tempat_wisata.xlsx'
    wb.save(response)
    return response

class PlaceListView(LoginRequiredMixin, ListView):
    model = TourismPlace
    template_name = 'tourism/place_list.html'
    context_object_name = 'places'
    paginate_by = 10  # pagination 10 item per halaman
    
    def get_queryset(self):
        # Urutkan berdasarkan yang terbaru
        return TourismPlace.objects.all().order_by('-created_at')


# View untuk Update/Edit Tempat Wisata
class PlaceUpdateView(LoginRequiredMixin, UpdateView):
    model = TourismPlace
    form_class = TourismPlaceForm
    template_name = 'tourism/edit_place.html'
    success_url = reverse_lazy('tourism_app:place_list')
    context_object_name = 'place'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # Kirim instance ke template jika perlu akses location
        context['location'] = self.object.location
        context['place'] = self.object  # supay template bisa akses objek
        return context
    
    def form_valid(self, form):
        messages.success(self.request, f'Data "{form.instance.name}" berhasil diperbarui!')
        return super().form_valid(form)
    
    def form_invalid(self, form):
        messages.error(self.request, 'Terjadi kesalahan. Periksa kembali data yang diisi.')
        return super().form_invalid(form)

# View untuk Delete/Hapus Tempat Wisata
class PlaceDeleteView(LoginRequiredMixin, DeleteView):
    model = TourismPlace
    template_name = 'tourism/place_confirm_delete.html'
    success_url = reverse_lazy('tourism_app:place_list')
    context_object_name = 'place'
    
    def delete(self, request, *args, **kwargs):
        place_name = self.get_object().name
        messages.success(request, f'Tempat wisata "{place_name}" berhasil dihapus!')
        return super().delete(request, *args, **kwargs)